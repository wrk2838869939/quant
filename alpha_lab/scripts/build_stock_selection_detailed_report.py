#!/usr/bin/env python3
"""把个股研究完整展开为图表、逐事件数据、Excel及可离线阅读的报告。"""
from pathlib import Path
import base64
import hashlib
import html
import json
import re
import shutil
import sys
import zipfile

import numpy as np
import pandas as pd
import yaml
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.colors import ListedColormap, BoundaryNorm, TwoSlopeNorm
from matplotlib.patches import Patch

ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT / 'alpha_lab/src'), str(ROOT / 'alpha_lab/scripts')]
from sector_dataset import digest
from report_stock_selection_asof_v2 import LABELS

RUN = ROOT / 'alpha_lab/outputs/stock_selection_asof_v2/20261013T074542_476130Z'
PARENT = ROOT / 'alpha_lab/outputs/sector_asof_replay_v1/20261013T072855_087132Z'
DOCS = ROOT / 'alpha_lab/docs/current_strategy_20261013'
OUT = DOCS / 'stock_selection_v2_detail'
FIG = OUT / 'figures'
TABLES = OUT / 'data'
ORDER = ['original', 'no_expansion', 'price_trend', 'trend_only', 'price_only',
         'price_strength5', 'price_lowvol5', 'ridge_price5', 'all_members']
COLORS = {'ridge': '#bf5a35', 'pool': '#367b9d', 'gray': '#74828d', 'green': '#40957e'}


def percent(x, signed=True):
    if pd.isna(x):
        return '—'
    return f'{x * 100:+.2f}%' if signed else f'{x * 100:.1f}%'


def table(headers, rows):
    return '\n'.join(['| ' + ' | '.join(headers) + ' |', '|' + '|'.join(['---'] * len(headers)) + '|']
                     + ['| ' + ' | '.join(str(x) for x in row) + ' |' for row in rows])


def savefig(fig, name):
    for artist in fig.findobj(matplotlib.text.Text):
        artist.set_text(artist.get_text().replace('\u2212', '-'))
    fig.savefig(FIG / (name + '.png'), dpi=180, bbox_inches='tight')
    fig.savefig(FIG / (name + '.svg'), bbox_inches='tight')
    plt.close(fig)


def load_data():
    """核对研究产物并读取完整固定时点样本。"""
    manifest = json.loads((RUN / 'manifest.json').read_text())
    assert manifest['status'] == 'complete'
    verified = 0
    for name, expected in manifest['artifacts'].items():
        assert digest(RUN / name) == expected, name
        verified += 1
    d = {}
    for key, name in [('summary', 'summary_NOT_ACCOUNT.csv'), ('annual', 'annual_NOT_ACCOUNT.csv'),
                      ('baskets', 'event_baskets_NOT_ACCOUNT.csv'), ('pairs', 'same_event_capped_allocation.csv'),
                      ('horizons', 'ridge_vs_price_pool_summary.csv'), ('year_pairs', 'ridge_same_pool_annual.csv'),
                      ('coefficients', 'model_coefficients.csv'), ('gates', 'gate_paired_comparisons.csv'),
                      ('turnover', 'turnover_gate_pairs.csv'), ('stress', 'missing_path_stress_NOT_ACCOUNT.csv'),
                      ('sensitivity', 'sensitivity.csv')]:
        d[key] = pd.read_csv(RUN / name)
    d['frame'] = pd.read_parquet(RUN / 'all_decisions_BEFORE_EVALUATION.parquet')
    d['fits'] = json.loads((RUN / 'fits.json').read_text())
    d['cfg'] = yaml.safe_load((RUN / 'config.yaml').read_text())
    d['parent_cfg'] = yaml.safe_load((PARENT / 'config.yaml').read_text())
    d['prices'] = pd.read_parquet(PARENT / 'observation_hfq_open.parquet')
    d['volume'] = pd.read_parquet(PARENT / 'observation_vol_share.parquet')
    d['observations'] = pd.read_parquet(PARENT / 'price_observations_NOT_ACCOUNT.parquet')
    catalog = pd.read_parquet(ROOT / d['parent_cfg']['dataset'] / 'catalog_sw_l2.parquet')
    d['names'] = catalog.set_index(['src', 'index_code']).industry_name.to_dict()
    d['verified'] = verified
    return d


def sector_name(d, code, date):
    version = 'SW2021' if pd.Timestamp(date) >= pd.Timestamp('2021-12-13') else 'SW2014'
    return d['names'].get((version, code), code)


def compute_paths(d):
    """重算每个排序事件的开盘价路径；不拼接事件为账户净值。"""
    frame, prices, volumes = d['frame'], d['prices'], d['volume']
    selected = frame[frame.ridge_price5]
    ids = selected.event_id.unique()
    paths, risks = [], []
    for event_id in ids:
        universe = frame[frame.event_id.eq(event_id)]
        date = universe.date.iloc[0]
        pos = prices.index.get_loc(date) + 1
        dates = prices.index[pos:pos + 21]
        for variant, basket in [('ridge_price5', universe[universe.ridge_price5]), ('price_only', universe[universe.price])]:
            codes = basket.ts_code.tolist()
            quote = prices.loc[dates, codes]
            volume = volumes.loc[dates, codes]
            complete = len(dates) == 21 and quote.gt(0).all().all() and volume.gt(0).all().all()
            if not complete:
                continue
            gross = quote.divide(quote.iloc[0], axis='columns').mean(axis=1)
            drawdown = gross / gross.cummax() - 1
            for offset, (day, value) in enumerate(gross.items()):
                paths.append(dict(event_id=event_id, date=date, trade_date=day, offset=offset,
                                  variant=variant, stocks=len(codes), gross_value=value,
                                  relative_to_entry=value - 1, drawdown=drawdown.iloc[offset]))
            risks.append(dict(event_id=event_id, date=date, variant=variant, stocks=len(codes),
                              min_relative_to_entry=float(gross.min() - 1), opening_path_mdd=float(drawdown.min()),
                              gross_end=float(gross.iloc[-1] - 1)))
    paths, risks = pd.DataFrame(paths), pd.DataFrame(risks)
    costs = d['parent_cfg']['observation']
    ratio = (1 - costs['sell_cost'] - costs['slippage_per_side']) / (1 + costs['buy_cost'] + costs['slippage_per_side'])
    stored = d['baskets'].query("scope == 'stock_model_available' and horizon == 20")
    checked = risks.merge(stored[['event_id', 'variant', 'price_return']], on=['event_id', 'variant'], validate='one_to_one')
    np.testing.assert_allclose((checked.gross_end + 1) * ratio - 1, checked.price_return, atol=1e-12, rtol=1e-12)
    paths.to_csv(TABLES / 'daily_open_paths.csv', index=False)
    risks.to_csv(TABLES / 'within_event_path_risk.csv', index=False)
    d['paths'], d['risks'] = paths, risks
    d['path_endpoint_checks'] = len(checked)


def figures(d):
    """所有分布、路径及案例均来自固定候选，不重新选策略。"""
    s20 = d['summary'].query('horizon == 20')
    all_s = s20.query("scope == 'all_sector_pass'").set_index('variant')
    common = s20.query("scope == 'stock_model_available'").set_index('variant')
    frame = d['frame']
    # 01：机会数量与数据完整性。
    names = ['original', 'no_expansion', 'price_trend', 'trend_only', 'price_only', 'price_strength5', 'price_lowvol5']
    fig, axes = plt.subplots(1, 2, figsize=(13, 5.3), layout='constrained')
    for ax, field, title in [(axes[0], 'signals', '确认日产生候选的事件 / 105'), (axes[1], 'stocks', '确认日候选：事件×股票数')]:
        vals = all_s.loc[names, field]
        ax.barh([LABELS[n] for n in names], vals, color=COLORS['pool'])
        for i, v in enumerate(vals):
            ax.text(v + max(vals) * .01, i, f'{int(v)}', va='center')
        ax.invert_yaxis(); ax.set_title(title); ax.set_xlim(0, max(vals) * 1.16); ax.grid(axis='x', alpha=.15)
    fig.suptitle('图1｜全部105个板块通过事件：数量变化', fontsize=15)
    savefig(fig, '01_opportunities')
    # 02：105事件逐条显示，模型不可用与空信号严格区分。
    events = frame[['event_id', 'date', 'sector_code']].drop_duplicates().sort_values(['date', 'event_id'])
    ids = events.event_id.tolist()
    matrix = np.full((len(ORDER), len(ids)), -1., dtype=float)
    for i, variant in enumerate(ORDER):
        scope = 'stock_model_available' if variant == 'ridge_price5' else 'all_sector_pass'
        rows = d['baskets'].query('horizon == 20')
        rows = rows[rows.scope.eq(scope) & rows.variant.eq(variant)].set_index('event_id')
        for j, event_id in enumerate(ids):
            if event_id in rows.index:
                row = rows.loc[event_id]
                matrix[i, j] = 0 if not row.signaled else 1 if row.complete else 2
    fig, ax = plt.subplots(figsize=(15, 5), layout='constrained')
    cmap = ListedColormap(['#c6ccd1', '#fafafa', '#3885a3', '#dc9564'])
    ax.imshow(matrix, cmap=cmap, norm=BoundaryNorm([-1.5, -.5, .5, 1.5, 2.5], 4), aspect='auto', interpolation='none')
    ax.set_yticks(range(len(ORDER)), [LABELS[x] for x in ORDER])
    ticks = np.arange(0, len(ids), 10)
    ax.set_xticks(ticks, [events.iloc[i].date.strftime('%Y-%m-%d') for i in ticks], rotation=35, ha='right', fontsize=9)
    for i in range(1, len(events)):
        if events.iloc[i].date.year != events.iloc[i - 1].date.year:
            ax.axvline(i - .5, color='#333', lw=.7)
    ax.legend(handles=[Patch(color=cmap(i), label=t) for i, t in enumerate(['模型历史不足', '当天无候选', '候选路径完整', '候选路径有缺失'])],
              loc='upper center', bbox_to_anchor=(.5, -.25), ncol=4, fontsize=10)
    ax.set_title('图2｜105个独立事件逐条覆盖；按日期及事件代码排列')
    savefig(fig, '02_event_coverage')
    # 03：共同可评分期的均值区间。
    names = [n for n in ORDER if n != 'all_members']
    fig, ax = plt.subplots(figsize=(11.7, 5.5), layout='constrained')
    for i, n in enumerate(names):
        row = common.loc[n]; color = COLORS['ridge'] if n == 'ridge_price5' else COLORS['pool']
        if np.isfinite(row.ci_low):
            ax.plot([row.ci_low * 100, row.ci_high * 100], [i, i], color=color, lw=2)
        ax.scatter(row.mean_excess * 100, i, color=color, s=45)
        if not np.isfinite(row.ci_low):
            ax.text(row.mean_excess * 100 + .1, i, '簇不足5，不报区间', va='center', fontsize=9)
    ax.axvline(0, color='#777', ls='--'); ax.invert_yaxis()
    ax.set_yticks(range(len(names)), [f'{LABELS[n]}（配对{int(common.loc[n, "paired_events"])}事件）' for n in names])
    ax.set(title='图3｜2022—2026同日板块基准：20日均值差及95%区间', xlabel='百分点；事件等权，行情簇重采样')
    ax.grid(axis='x', alpha=.15); savefig(fig, '03_common_period_intervals')
    # 04：年度绝对观察值与相对值；此处不将不同事件覆盖视为控制实验。
    years = [2022, 2023, 2024, 2025, 2026]
    annual = d['annual'].query("scope == 'stock_model_available' and horizon == 20")
    names = [n for n in ORDER if n != 'all_members']
    fig, axes = plt.subplots(1, 2, figsize=(14, 7), layout='constrained')
    for ax, metric, title in [(axes[0], 'mean', '候选篮子20日均值'), (axes[1], 'mean_excess', '相对同日板块成员均值')]:
        vals = annual.pivot(index='variant', columns='year', values=metric).reindex(index=names, columns=years) * 100
        limit = max(1., np.nanmax(np.abs(vals)))
        im = ax.imshow(vals, cmap='RdBu', norm=TwoSlopeNorm(vmin=-limit, vcenter=0, vmax=limit), aspect='auto')
        for i in range(len(names)):
            for j in range(len(years)):
                value = vals.iloc[i, j]
                ax.text(j, i, '—' if pd.isna(value) else f'{value:+.1f}', ha='center', va='center',
                        color='white' if abs(value) > limit * .65 else '#222', fontsize=10)
        ax.set_xticks(range(len(years)), years); ax.set_yticks(range(len(names)), [LABELS[n] for n in names]); ax.set_title(title)
        fig.colorbar(im, ax=ax, shrink=.7, label='百分比（%）' if metric == 'mean' else '百分点')
    fig.suptitle('图4｜年度差异：各方案有信号的事件均值（数量详见表格）')
    savefig(fig, '04_yearly_heatmaps')
    # 05：52个相同事件中的分布。
    pair = d['pairs']
    diff = (pair.price_return_ml - pair.price_return_base) * 100
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.8), layout='constrained')
    for field, label, color in [('price_return_ml', '历史训练排序', COLORS['ridge']), ('price_return_base', '同日价格形态池', COLORS['pool'])]:
        x = np.sort(pair[field].to_numpy() * 100)
        axes[0].step(x, np.arange(1, len(x) + 1) / len(x), where='post', color=color, label=label, lw=2)
    axes[0].set(xlabel='20日费用压力测试后的价格观察（%）', ylabel='累计比例', title='同一52事件的经验分布')
    axes[0].legend(); axes[0].grid(alpha=.15)
    axes[1].hist(diff, bins=18, color=COLORS['green'], edgecolor='white')
    axes[1].axvline(0, color='#555', ls='--'); axes[1].axvline(diff.mean(), color=COLORS['ridge'], lw=2, label=f'均值{diff.mean():+.2f}个百分点')
    axes[1].set(xlabel='排序减价格池（百分点）', ylabel='事件数', title='每事件配对差值'); axes[1].legend()
    fig.suptitle('图5｜同事件同日期分布，不按单只股票扩充样本数')
    savefig(fig, '05_matched_distributions')
    # 06：交易后的事件内开盘价格路径，仅统计同一配对样本。
    paths = d['paths'][d['paths'].event_id.isin(pair.event_id)]
    fig, axes = plt.subplots(1, 2, figsize=(12, 5), layout='constrained')
    for variant, label, color in [('ridge_price5', '历史训练排序', COLORS['ridge']), ('price_only', '价格形态池', COLORS['pool'])]:
        group = paths[paths.variant.eq(variant)].groupby('offset').relative_to_entry
        mean, median, low, high = group.mean() * 100, group.median() * 100, group.quantile(.25) * 100, group.quantile(.75) * 100
        axes[0].plot(mean.index, mean, color=color, label=label, lw=2)
        axes[1].plot(median.index, median, color=color, label=label, lw=2)
        axes[1].fill_between(low.index, low, high, color=color, alpha=.12)
    for ax in axes:
        ax.axhline(0, color='#777', lw=.8); ax.set_xlabel('相对买入开盘的交易日偏移'); ax.set_ylabel('相对买入价（%）'); ax.grid(alpha=.15); ax.legend()
    axes[0].set_title('事件等权均值路径'); axes[1].set_title('中位路径与25%—75%分位带')
    fig.suptitle('图6｜52配对事件的复权开盘价格路径；未扣费，不是账户净值')
    savefig(fig, '06_open_price_paths')
    # 07：事件内不利路径，不能命名为策略账户最大回撤。
    risks = d['risks'][d['risks'].event_id.isin(pair.event_id)]
    fig, axes = plt.subplots(1, 2, figsize=(12, 5), layout='constrained')
    riskrows = []
    for ax, field, title in [(axes[0], 'min_relative_to_entry', '期间最差开盘价相对买入价'),
                              (axes[1], 'opening_path_mdd', '事件内开盘价格路径最大回撤')]:
        data = [risks.loc[risks.variant.eq(v), field].to_numpy() * 100 for v in ['ridge_price5', 'price_only']]
        boxes = ax.boxplot(data, tick_labels=['历史训练排序', '价格形态池'], patch_artist=True, showmeans=True)
        for patch, color in zip(boxes['boxes'], [COLORS['ridge'], COLORS['pool']]):
            patch.set_facecolor(color); patch.set_alpha(.5)
        ax.set(title=title, ylabel='百分比'); ax.axhline(0, color='#888', lw=.8); ax.grid(axis='y', alpha=.15)
        for variant in ['ridge_price5', 'price_only']:
            values = risks.loc[risks.variant.eq(variant), field]
            riskrows.append(dict(metric=field, variant=variant, events=len(values), mean=values.mean(), median=values.median(), q10=values.quantile(.1), worst=values.min()))
    fig.suptitle('图7｜52配对事件的路径风险；不包含盘中低点与账户资金重叠')
    savefig(fig, '07_within_event_risk')
    d['risk_summary'] = pd.DataFrame(riskrows)
    d['risk_summary'].to_csv(TABLES / 'path_risk_summary.csv', index=False)
    # 08：持有期敏感性。
    h = d['horizons']
    fig, ax = plt.subplots(figsize=(10, 4.6), layout='constrained')
    for i, row in enumerate(h.itertuples()):
        ax.plot([row.ci_low * 100, row.ci_high * 100], [i, i], color=COLORS['pool'], lw=3)
        ax.scatter(row.mean * 100, i, color=COLORS['ridge'], s=65)
    ax.set_yticks(range(len(h)), [f'{r.horizon}交易日（{r.pairs}配对事件）' for r in h.itertuples()]); ax.invert_yaxis()
    ax.axvline(0, color='#777', ls='--'); ax.grid(axis='x', alpha=.15)
    ax.set(title='图8｜排序相对同日价格池：观察期变化', xlabel='均值差及95%区间（百分点）')
    savefig(fig, '08_horizon_sensitivity')
    # 09：仓位集中与留现金的效应。
    counts = frame[frame.ridge_price5].groupby('event_id').size().value_counts().reindex(range(1, 6), fill_value=0)
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.7), layout='constrained')
    axes[0].bar(counts.index, counts, color=COLORS['pool']); axes[0].set_xticks(range(1, 6))
    for x, y in counts.items():
        axes[0].text(x, y + .4, str(y), ha='center')
    axes[0].set(title='54个排序事件选出多少只股票', xlabel='入选股票数', ylabel='事件数', ylim=(0, counts.max() * 1.2))
    labels = ['入选篮子用满额度', '每股≤20%，剩余现金']
    values = [(pair.price_return_ml - pair.price_return_base).mean(), (pair.capped_return_ml - pair.capped_return_base).mean()]
    axes[1].bar(labels, np.array(values) * 100, color=[COLORS['ridge'], COLORS['pool']])
    for i, v in enumerate(values):
        axes[1].text(i, v * 100 + .035, f'{v * 100:+.2f}', ha='center')
    axes[1].set(title='同一52事件的均值差变化', ylabel='排序减价格池（百分点）', ylim=(0, max(values) * 120))
    fig.suptitle('图9｜候选收益不能替代仓位收益；事件额度仍非账户')
    savefig(fig, '09_allocation_effect')
    # 10：模型随年变化。
    translations = {'strength_rank': '60日强度分位', 'volatility_rank': '低波动分位', 'expansion_rank': '前期放量分位',
                    'contraction_rank': '后期缩量分位', 'slowdown_rank': '跌速减缓分位', 'trend_flag': '趋势通过', 'price_flag': '价格形态通过',
                    'trend_price': '趋势×价格交互', 'strength_missing': '强度缺失', 'volatility_missing': '波动缺失',
                    'expansion_missing': '放量缺失', 'contraction_missing': '缩量缺失', 'slowdown_missing': '跌速缺失'}
    coef = d['coefficients'].pivot(index='feature', columns='year', values='coefficient') * 100
    fig, ax = plt.subplots(figsize=(10, 7), layout='constrained')
    limit = np.abs(coef.to_numpy()).max()
    im = ax.imshow(coef, cmap='RdBu', norm=TwoSlopeNorm(vmin=-limit, vcenter=0, vmax=limit), aspect='auto')
    for i in range(len(coef)):
        for j in range(len(coef.columns)):
            v = coef.iloc[i, j]; ax.text(j, i, f'{v:+.2f}', ha='center', va='center', color='white' if abs(v) > limit * .65 else '#222', fontsize=10)
    ax.set_yticks(range(len(coef)), [translations[k] for k in coef.index]); ax.set_xticks(range(len(coef.columns)), coef.columns)
    ax.set_title('图10｜每年历史训练的岭回归系数；不是因果效应')
    fig.colorbar(im, ax=ax, label='每单位去均值特征对应的预测差（百分点）')
    savefig(fig, '10_model_coefficients')
    # 11：门槛通过与未通过组的配对比较。
    gs = d['gates'].copy()
    labels = ['趋势＋价格 - 仅价格', '趋势＋价格 - 仅趋势', '仅价格 - 二者都未过', '仅趋势 - 二者都未过']
    ts = d['turnover'].rename(columns={'gate': 'comparison'})
    allg = pd.concat([gs, ts], ignore_index=True)
    labels += ['趋势价格池：放量通过 - 未通过', '趋势价格池：缩量通过 - 未通过']
    fig, ax = plt.subplots(figsize=(12, 5.3), layout='constrained')
    for i, row in enumerate(allg.itertuples()):
        ax.plot([row.ci_low * 100, row.ci_high * 100], [i, i], color=COLORS['pool'], lw=2)
        ax.scatter(row.mean * 100, i, color=COLORS['ridge'], s=45)
    ax.set_yticks(range(len(labels)), [f'{t}（n={int(n)}）' for t, n in zip(labels, allg.paired_events)])
    ax.invert_yaxis(); ax.axvline(0, color='#777', ls='--'); ax.grid(axis='x', alpha=.15)
    ax.set(title='图11｜同一事件同一天：个股门槛的边际差异', xlabel='20日均值差及95%区间（百分点）')
    savefig(fig, '11_gate_comparisons')
    # 12：按明确规则挑选案例，不以好坏决定是否纳入总体。
    pair = d['pairs'].copy()
    pair['difference'] = pair.price_return_ml - pair.price_return_base
    candidates = [('最大相对改善', pair.sort_values('difference', ascending=False)),
                  ('最大相对落后', pair.sort_values('difference')),
                  ('最接近配对差中位数', pair.assign(distance=(pair.difference - pair.difference.median()).abs()).sort_values('distance')),
                  ('排序绝对表现最差', pair.sort_values('price_return_ml')),
                  ('最近事件之一', pair.sort_values(['date_ml', 'event_id'], ascending=False)),
                  ('最近另一事件', pair.sort_values(['date_ml', 'event_id'], ascending=False))]
    cases, used = [], set()
    for reason, ordered in candidates:
        row = ordered[~ordered.event_id.isin(used)].iloc[0]; used.add(row.event_id)
        subset = frame[frame.event_id.eq(row.event_id)]
        code = subset.sector_code.iloc[0]
        cases.append(dict(reason=reason, event_id=row.event_id, date=row.date_ml, sector_code=code,
                          sector_name=sector_name(d, code, row.date_ml), stocks=int(row.stocks_ml),
                          stock_codes=','.join(subset.loc[subset.ridge_price5, 'ts_code']), model_return=row.price_return_ml,
                          pool_return=row.price_return_base, difference=row.difference))
    cases = pd.DataFrame(cases); d['cases'] = cases; cases.to_csv(TABLES / 'case_selection_and_results.csv', index=False)
    fig, axes = plt.subplots(3, 2, figsize=(14, 12), layout='constrained')
    for ax, row in zip(axes.flat, cases.itertuples()):
        for variant, label, color in [('ridge_price5', '历史训练排序', COLORS['ridge']), ('price_only', '价格形态池', COLORS['pool'])]:
            path = d['paths'][d['paths'].event_id.eq(row.event_id) & d['paths'].variant.eq(variant)]
            ax.plot(path.offset, path.gross_value * 100, label=label, color=color, lw=2)
        ax.axhline(100, color='#888', lw=.8); ax.grid(alpha=.15)
        ax.set(title=f'{row.date} {row.sector_name}\n{row.reason}｜20日差{row.difference * 100:+.2f}个百分点',
               xlabel='买入开盘后的交易日偏移', ylabel='复权开盘价格（买入=100）')
        ax.legend(fontsize=9)
    fig.suptitle('图12｜诊断案例：按事后好坏及最近日期选取，仅作解释，不代表总体', fontsize=14)
    savefig(fig, '12_cases')
    d['event_order'] = events


def export_tables(d):
    """提供完整105事件、193候选和3935股票诊断，不只导出汇总。"""
    frame = d['frame']
    rules = [n for n in ORDER if n != 'all_members']
    event = frame.groupby(['event_id', 'date', 'sector_code'])[rules].sum().reset_index()
    event['sector_name'] = [sector_name(d, r.sector_code, r.date) for r in event.itertuples()]
    available = frame.groupby('event_id').model_start.apply(lambda x: x.notna().any())
    event['stock_model_available'] = event.event_id.map(available)
    model_return = d['baskets'].query("scope == 'stock_model_available' and variant == 'ridge_price5' and horizon == 20").set_index('event_id').price_return
    event['model_20d_price_return'] = event.event_id.map(model_return)
    event = event.sort_values(['date', 'event_id'])
    obs = d['observations'].merge(d['event_order'][['event_id', 'date']], on=['event_id', 'date'], validate='many_to_one')
    selected = frame[frame.ridge_price5].copy()
    for horizon in [3, 10, 20]:
        subset = obs[obs.horizon.eq(horizon)][['event_id', 'ts_code', 'price_return', 'complete', 'entry_date', 'exit_date']]
        subset = subset.rename(columns={c: f'{c}_{horizon}d' for c in subset.columns if c not in ['event_id', 'ts_code']})
        selected = selected.merge(subset, on=['event_id', 'ts_code'], validate='one_to_one')
    selected['sector_name'] = [sector_name(d, r.sector_code, r.date) for r in selected.itertuples()]
    training = []
    for r in d['fits']:
        if r['status'] == 'skipped':
            training.append(dict(year=r['year'], status=r['status'], reason=r['reason']))
        else:
            training.append(dict(year=r['year'], status=r['status'], train_events=len(r['train_events']),
                                 fit_events=len(r['fit_events']), valid_events=len(r['valid_events']), train_rows=r['train_rows'],
                                 alpha=r['alpha'], cutoff=r['cutoff'], train_last_label=r['train_last_label'],
                                 inner_cutoff=r['inner_cutoff'], fit_last_label=r['fit_last_label'], prefix_exact=r['prefix_exact']))
    training = pd.DataFrame(training); d['training'] = training; d['event_table'] = event; d['selected_table'] = selected
    dictionary = pd.DataFrame([
        ('event_id', '独立板块事件编号；不是股票交易笔数', '字符串'),
        ('scope', 'all_sector_pass=105事件；stock_model_available=2022—2026的59事件', '分组'),
        ('horizon', '次日开盘进入，经过指定交易日数量后开盘观察', '交易日'),
        ('price_return / mean / median / q10', '固定约40bp费用压力测试后的复权价格观察，不是账户收益', '小数；0.01=1%'),
        ('mean_excess / difference / excess', '相对同日同事件基准的差值，基准见各表', '小数；0.01=1个百分点'),
        ('positive', '完整价格路径事件中观察值大于0的比例', '小数；0.5=50%'),
        ('complete', '整个观察窗口开盘价和成交量都有效；缺失不能解释为零损失', '布尔'),
        ('capped_return', '每股最高20%事件额度、剩余现金为零收益；不处理事件之间资金冲突', '小数'),
        ('mean_slot / known_slots', '无候选事件为0；有候选但价格未知为缺失；均值分母为known_slots', '小数/事件数'),
        ('stress_20 / stress_100', '缺失股票按−20%或−100%替代的假设情景；不是实际损失或置信下界', '小数'),
        ('model_score', '历史岭回归预测的相对板块表现；不是保证收益', '小数'),
        ('volatility', '近20个交易日的日收益标准差', '小数'),
        ('strength', '截至确认日的60交易日累计价格变化', '小数'),
        ('expansion / contraction / slowdown', '前期放量倍数、后期换手比值、跌速比值', '倍数'),
        ('gross_value', '从买入开盘起，事件内初始等权持有的复权开盘价值指数；未扣费', '买入=1'),
        ('opening_path_mdd', '单个事件21个开盘观测点从先前最高点的最大下降', '小数；不是账户MDD'),
        ('min_relative_to_entry', '事件内最低开盘价值相对买入价的变化，不含盘中低点', '小数'),
        ('date / model_start', '信号收盘日 / 当年模型训练起点；不能据当前下载版本认证原始PIT', '日期'),
    ], columns=['字段', '说明', '单位'])
    tables = {
        '01_字段与口径': dictionary, '02_105事件总览': event, '03_完整统计': d['summary'],
        '04_全部年度': d['annual'], '05_全部事件篮子': d['baskets'], '06_193候选与结果': selected,
        '07_3935个股诊断': frame, '08_历史训练记录': training, '09_模型系数': d['coefficients'],
        '10_52同日配对': d['pairs'], '11_配对年度': d['year_pairs'], '12_3种观察期': d['horizons'],
        '13_事件内价格路径': d['paths'], '14_事件内风险': d['risks'], '15_路径风险汇总': d['risk_summary'],
        '16_条件配对': d['gates'], '17_换手配对': d['turnover'], '18_敏感性': d['sensitivity'],
        '19_缺失压力情景': d['stress'], '20_案例选择与结果': d['cases'],
    }
    with pd.ExcelWriter(OUT / 'DATA.xlsx', engine='openpyxl') as writer:
        for name, data in tables.items():
            data.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]; ws.freeze_panes = 'D2' if len(data.columns) > 6 else 'A2'
            ws.auto_filter.ref = ws.dimensions
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='244B64')
            for i, namecol in enumerate(data.columns, 1):
                ws.column_dimensions[get_column_letter(i)].width = min(38, max(15, len(str(namecol)) + 3))
            data.to_csv(TABLES / (name + '.csv'), index=False, encoding='utf-8-sig')
    d['sheets'] = {name: len(data) for name, data in tables.items()}


def write_report(d):
    """报告正文直接嵌入全部图表及细表，不以输出目录链接替代内容。"""
    s20 = d['summary'].query('horizon == 20')
    all_s = s20.query("scope == 'all_sector_pass'").set_index('variant')
    common = s20.query("scope == 'stock_model_available'").set_index('variant')
    pair = d['pairs']
    lines = ['# 板块模型通过后的个股筛选：详细数据与图表报告', '',
             '**研究日期：2026-10-13｜行情截至：2026-10-12｜完整版本：12组图、20张Excel数据表、105事件与193候选明细**', '',
             '本报告使用按历史时点训练的板块与个股模型。所有数字均为历史条件回放，价格观察施加固定费用压力测试；没有真实账户成交认证。供应商原始历史版本和行业公告时间仍有缺口。', '',
             '[下载完整Excel数据（20张表）](DATA.xlsx) · [PDF版](REPORT.pdf) · [自包含HTML版](REPORT.html) · [离线完整包](stock_selection_detail.zip)', '',
             '## 1. 结论、样本与比较口径', '',
             '**候选数量能显著增加，但稳定选股优势尚未确认。** 固定确认日后，原规则在105个板块事件中只有9次有候选；放宽前期放量为40次，趋势加价格为69次，价格形态池为103次。历史训练排序在2022—2026可评分的59个事件中产生54次候选。', '',
             '同一52事件内，排序相对价格池的20日均值差为+1.20个百分点，95%区间为−0.09～+2.05个百分点；每股最多20%、不足5只留现金后，差值缩小到+0.19个百分点。', '',
             table(['项目', '范围与定义'], [
                 ['板块样本', '2019—2026年105个按当时模型通过的独立事件'],
                 ['个股样本', '确认日3935个事件×股票；同一股票可属于不同历史事件'],
                 ['主入场口径', '板块首次确认日收盘选股，次日开盘观察；不等待后4日'],
                 ['观察期限', '3、10、20交易日；20日为预先指定主期限'],
                 ['费用', '买10bp、卖20bp、双边各5bp滑点；未模拟整手最低费用'],
                 ['原规则9与旧报告17', '9为确认当日；17为允许5日内等待，二者不能混用'],
                 ['统计单位', '板块事件等权；不把同一事件的多只股票当作独立样本'],
                 ['区间', '完整事件时间线先固定20交易日行情簇，再整簇重采样2000次'],
                 ['缺失处理', '任何篮子成分路径缺失，该篮子主表不报收益；缺失和压力情景保留'],
             ]), '',
             '## 2. 候选数量和逐事件覆盖', '', '![图1](figures/01_opportunities.png)', '',
             table(['规则', '有候选/105', '候选股票数', '完整事件', '可与板块配对事件'],
                   [[LABELS[n], int(all_s.loc[n, 'signals']), int(all_s.loc[n, 'stocks']),
                     int(all_s.loc[n, 'complete_signals']), int(all_s.loc[n, 'paired_events'])]
                    for n in ORDER if n not in ['ridge_price5']]), '',
             '![图2](figures/02_event_coverage.png)', '',
             '图2每列是一个独立事件：蓝色有候选且路径完整，橙色有候选但路径缺失，白色无候选，灰色为个股模型历史不足。最下方全体成员作为覆盖对照。模型缺历史不能算成主动空仓。', '',
             '## 3. 全部105事件：20日完整统计', '',
             table(['规则', '完整/有信号', '均值', '中位', '10%分位', '正值率', '相对板块'],
                   [[LABELS[n], f'{int(all_s.loc[n,"complete_signals"])}/{int(all_s.loc[n,"signals"])}',
                     percent(all_s.loc[n,'mean']), percent(all_s.loc[n,'median']), percent(all_s.loc[n,'q10']),
                     percent(all_s.loc[n,'positive'], False), percent(all_s.loc[n,'mean_excess'])]
                    for n in ORDER if n != 'ridge_price5']), '',
             '“相对板块”百分数读作百分点。相对值要求当日全体合格板块成员的路径也完整，因此配对数可能少于候选完整数。各规则覆盖事件不同，不能仅按表中均值挑赢家。', '',
             '## 4. 共同可评分期：2022—2026年59事件', '', '![图3](figures/03_common_period_intervals.png)', '',
             table(['规则', '有候选/59', '配对数', '20日均值', '相对板块', '相对值95%区间'],
                   [[LABELS[n], int(common.loc[n,'signals']), int(common.loc[n,'paired_events']), percent(common.loc[n,'mean']),
                     percent(common.loc[n,'mean_excess']), f'[{percent(common.loc[n,"ci_low"])}, {percent(common.loc[n,"ci_high"])}]'] for n in ORDER]), '',
             '除对照自身差值为零外，没有方案的相对板块区间排除零。原规则只4个可配对事件、3个行情簇，因此不报告其区间。', '',
             '## 5. 年度表现与样本量', '', '![图4](figures/04_yearly_heatmaps.png)', '',
             '下表每格为“20日均值 / 有信号事件数”。不同规则每年仍可能覆盖不同事件，真正同事件比较见下一节。', '']
    annual = d['annual'].query("scope == 'stock_model_available' and horizon == 20").set_index(['variant','year'])
    lines += [table(['规则', '2022', '2023', '2024', '2025', '2026'],
                    [[LABELS[n]] + [f'{percent(annual.loc[(n,y),"mean"])} / {int(annual.loc[(n,y),"signals"])}' for y in range(2022, 2027)] for n in ORDER]), '',
              '## 6. 同一52事件：排序与价格池的分布比较', '', '![图5](figures/05_matched_distributions.png)', '',
              table(['同一52事件', '均值', '中位', '10%分位', '最差', '最好'],
                    [[label, percent(pair[field].mean()), percent(pair[field].median()), percent(pair[field].quantile(.1)), percent(pair[field].min()), percent(pair[field].max())]
                     for label, field in [('排序', 'price_return_ml'), ('价格形态池', 'price_return_base')]]), '',
              table(['年份', '严格配对事件', '排序−价格池均值'],
                    [[int(r.date), int(r.events), percent(r.difference)] for r in d['year_pairs'].itertuples()]), '',
              '五个年度的配对差值方向均为正，但样本只有52个事件和9个行情簇。删除2022年后为+0.83个百分点，删除最大单事件贡献后为+0.85个百分点。与全体板块成员的比较是另一种基准，不能混用。', '',
              '## 7. 持有期间的开盘价格路径与风险', '', '![图6](figures/06_open_price_paths.png)', '',
              '每个事件从买入开盘设为1，股票初始等权买入并持有，再对事件做均值或分位统计。图中未扣费，不把不同日期事件拼成账户净值。分位带表示事件差异，不是预测置信区间。', '',
              '![图7](figures/07_within_event_risk.png)', '',
              table(['路径指标', '篮子', '事件数', '均值', '中位', '10%分位', '最差'],
                    [[('相对买入价最差开盘' if r.metric == 'min_relative_to_entry' else '事件内开盘路径最大回撤'), LABELS[r.variant],
                      r.events, percent(r.mean), percent(r.median), percent(r.q10), percent(r.worst)] for r in d['risk_summary'].itertuples()]), '',
              '**这些数值不是账户最大回撤。** 开盘路径没有盘中最低价、限价成交、停牌卖不出、事件间资金重叠及公司行动现金账本，不能据此声称满足账户回撤目标。', '',
              '## 8. 观察期限与仓位约束', '', '![图8](figures/08_horizon_sensitivity.png)', '',
              table(['观察期', '配对事件', '差值均值', '差值中位', '95%区间'],
                    [[r.horizon, r.pairs, percent(r.mean), percent(r.median), f'[{percent(r.ci_low)}, {percent(r.ci_high)}]'] for r in d['horizons'].itertuples()]), '',
              '3日和10日没有一致改善；不能只展示20日正值后声称所有持有期限都有效。', '',
              '![图9](figures/09_allocation_effect.png)', '',
              table(['同一52事件', '排序', '价格形态池', '差值'], [
                  ['平均投入比例', percent(pair.capped_exposure_ml.mean(),False), percent(pair.capped_exposure_base.mean(),False), '—'],
                  ['用满事件额度的20日观察', percent(pair.price_return_ml.mean()), percent(pair.price_return_base.mean()), percent((pair.price_return_ml-pair.price_return_base).mean())],
                  ['每股≤20%、剩余现金', percent(pair.capped_return_ml.mean()), percent(pair.capped_return_base.mean()), percent((pair.capped_return_ml-pair.capped_return_base).mean())],
              ]), '',
              '仓位约束后的差值95%区间为−1.21～+1.84个百分点。这里的每个事件额度独立，仍不是可相加的账户收益。不能把1只候选默认占满账户。', '',
              '## 9. 模型的历史训练记录与系数', '',
              table(['年份', '训练事件', '内层拟合', '内层验证', '正则α', '训练最晚标签', '状态'],
                    [[int(r.year), '—' if pd.isna(r.train_events) else int(r.train_events),
                      '—' if pd.isna(r.fit_events) else int(r.fit_events), '—' if pd.isna(r.valid_events) else int(r.valid_events),
                      '—' if pd.isna(r.alpha) else int(r.alpha), str(r.train_last_label)[:10] if pd.notna(r.train_last_label) else '—',
                      '历史不足跳过' if r.status == 'skipped' else '删除未来后权重完全一致'] for r in d['training'].itertuples()]), '',
              '每年模型训练及内层验证都按日期拆分，标签须在起点前20交易日截止之前结束。正则参数在历史内层选择；当年的特征只用于预测，不参与训练和选参。训练日期等权，同日事件等权，事件内股票等权。', '',
              '![图10](figures/10_model_coefficients.png)', '',
              '系数对应去均值后的截面特征，不是因果解释；分位方向中“低波动分位”越高表示越低波。模型的低波分位系数为负，有偏好高波动股票的迹象，不能仅用均值改善忽略风险。缺失标记也是当时可见的信息。', '',
              '## 10. 单个门槛的同日配对证据', '', '![图11](figures/11_gate_comparisons.png)', '',
              table(['比较', '配对事件', '20日均值差', '95%区间'],
                    [[label, int(r.paired_events), percent(r.mean), f'[{percent(r.ci_low)}, {percent(r.ci_high)}]']
                     for label,r in zip(['趋势＋价格 对 仅价格','趋势＋价格 对 仅趋势','仅价格 对 两者均未过','仅趋势 对 两者均未过'],d['gates'].itertuples())]
                    + [[('趋势价格池内：前期放量通过' if r.gate=='expansion' else '趋势价格池内：后期缩量通过'), int(r.paired_events), percent(r.mean), f'[{percent(r.ci_low)}, {percent(r.ci_high)}]'] for r in d['turnover'].itertuples()]), '',
              '这里比较的是相同板块事件同一天的不同个股组。门槛通过组方向偏正，但区间较宽，不能证明必须一律硬过滤，也不能由不显著推出门槛完全无效。', '',
              '## 11. 六个诊断案例：包括落后和亏损', '', '![图12](figures/12_cases.png)', '',
              table(['日期/板块', '选例规则', '排序/价格池20日', '差值', '入选代码'],
                    [[f'{r.date} {r.sector_name}', r.reason, f'{percent(r.model_return)} / {percent(r.pool_return)}', percent(r.difference), r.stock_codes] for r in d['cases'].itertuples()]), '',
              '案例用已实现结果选择，只用于解释，不参与训练或总体筛选。最大改善、最大落后、接近中位、最差绝对表现和最近事件全部披露，避免以几个漂亮曲线代表总体。价格图是复权开盘未扣费路径，表格是扣固定费用后的终点值。', '',
              '## 12. 缺失数据、压力情景与可交易性', '',
              '主表对不完整篮子不报收益，保留原候选。下面将缺失股票统一替换为−20%/−100%，其余已知股票照常计，且无候选事件为0；这是人为情景，不是实际损失估计或可靠下界。情景会纳入主表排除的整个事件，因此不保证严格低于主表均值。', '',
              table(['2022—2026，20日', '已知/全部事件额度', '主表事件额度均值', '缺失−20%情景', '缺失−100%情景'],
                    [[LABELS[n], f'{int(common.loc[n,"known_slots"])}/59', percent(common.loc[n,'mean_slot']),
                      percent(d['stress'].query("scope == 'stock_model_available' and horizon == 20").set_index('variant').loc[n,'stress_20']),
                      percent(d['stress'].query("scope == 'stock_model_available' and horizon == 20").set_index('variant').loc[n,'stress_100'])] for n in ORDER]), '',
              '原始历史行业区间缺少公告/接收时间，日线和复权数据也可能被供应商修订；本轮规则家族由今天制定。当前只能核验算法时间边界，不能认证当年实际收到的数据版本。官方涨跌停覆盖、排队成交、整手、公司行动到账和统一资金占用尚未完成验证。', '',
              '## 13. 105个事件逐条明细', '',
              '各列是确认当日候选股票数。排序“未训练”表示过去样本不足；0表示模型可用但未选股。最右列为排序篮子的20日费用压力测试价格观察，空值不等于0。完整193个候选、3935个股诊断和每日路径见Excel与data目录。', '',
              table(['日期', '板块', '原规则', '去放量', '趋势价格', '价格池', '排序', '排序20日'],
                    [[r.date.strftime('%Y-%m-%d'), r.sector_name, r.original, r.no_expansion, r.price_trend, r.price_only,
                      r.ridge_price5 if r.stock_model_available else '未训练', percent(r.model_20d_price_return)] for r in d['event_table'].itertuples()]), '',
              '## 14. 数据下载、复现与核验', '',
              '[完整Excel：DATA.xlsx](DATA.xlsx) 包含以下20张表；数值单位见第一张表，所有原始收益小数0.01代表1%。', '',
              table(['数据表', '行数'], list(d['sheets'].items())), '',
              f'本次核验研究产物{d["verified"]}个哈希；新增开盘路径的{d["path_endpoint_checks"]}个篮子终点与原研究结果逐值一致。历史时序专项测试3项通过，5个真实年度删除未来后重训通过，主研究审计47项哈希通过。', '',
              '复现命令：`python alpha_lab/scripts/build_stock_selection_detailed_report.py`。生成HTML、PDF、Markdown、PNG/SVG、Excel和CSV，不改变模型、候选或生产配置。报告生成脚本及数据源哈希见detail_manifest.json。', '']
    text = '\n'.join(lines).replace('\u2212', '-')
    (OUT / 'REPORT.md').write_text(text)
    primary = text.replace('](figures/', '](stock_selection_v2_detail/figures/').replace('](DATA.xlsx)', '](stock_selection_v2_detail/DATA.xlsx)')
    primary = primary.replace('](REPORT.pdf)', '](stock_selection_v2_detail/REPORT.pdf)').replace('](REPORT.html)', '](stock_selection_v2_detail/REPORT.html)').replace('](stock_selection_detail.zip)', '](stock_selection_v2_detail/stock_selection_detail.zip)')
    (DOCS / 'STOCK_SELECTION_V2_RESULTS.md').write_text(primary)
    return text


def render_and_package(d, text):
    """离线HTML嵌入图片，PDF逐页检查可读文字与图像，打包完整数据。"""
    import markdown
    import pymupdf
    from PIL import Image
    body = markdown.markdown(text, extensions=['tables', 'fenced_code'])
    # HTML版图片全部内嵌，单独打开也能看到图，不依赖输出目录。
    def inline_image(match):
        path = OUT / match.group(1)
        encoded = base64.b64encode(path.read_bytes()).decode()
        return f'src="data:image/png;base64,{encoded}"'
    html_body = re.sub(r'src="(figures/[^\"]+\.png)"', inline_image, body)
    web_css = '''body{font-family:system-ui,"Microsoft Yahei",sans-serif;max-width:1320px;margin:auto;padding:30px;color:#233947;line-height:1.65;background:#fff}h1,h2{color:#173c54}h2{border-top:2px solid #d9e5ec;padding-top:22px;margin-top:42px}table{border-collapse:collapse;width:100%;font-size:13px;margin:20px 0}td,th{border:1px solid #cad7df;padding:8px;text-align:left}th{background:#e9f1f5;position:sticky;top:0}tr:nth-child(even){background:#f7fafc}img{width:100%;height:auto;border:1px solid #dde5ea}a{color:#176794}code{overflow-wrap:anywhere}@media print{body{max-width:none;padding:0}h2{break-before:page}}'''
    (OUT / 'REPORT.html').write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>个股筛选详细研究</title><style>' + web_css + '</style><body>' + html_body + '</body></html>')
    css = '''@font-face{font-family:ReportCN;src:url(SimHei.ttf)}body{font-family:ReportCN,sans-serif;font-size:9.2pt;line-height:1.45;color:#203744}h1{font-size:20pt;color:#163d54}h2{font-size:14pt;color:#163d54;page-break-before:always;page-break-after:avoid;margin:12pt 0}p{margin:6pt 0}table{border-collapse:collapse;width:100%;font-size:7.2pt;margin:9pt 0}td,th{border:.5pt solid #c2d0d9;padding:4pt;vertical-align:top}th{background-color:#e7eff4}img{width:100%}a{color:#225f83}code{font-family:ReportCN,sans-serif;font-size:8pt}'''
    archive = pymupdf.Archive(str(OUT)); archive.add(str(ROOT / 'fonts'))
    story = pymupdf.Story(html=body, user_css=css, archive=archive)
    pdf_temp = OUT / 'layout_work.pdf'
    writer = pymupdf.DocumentWriter(str(pdf_temp))
    page = pymupdf.paper_rect('a4'); area = pymupdf.Rect(34, 36, page.width - 34, page.height - 34)
    pages = 0
    while True:
        device = writer.begin_page(page); more, _ = story.place(area); story.draw(device); writer.end_page(); pages += 1
        if not more:
            break
        if pages > 70:
            raise RuntimeError('PDF排版超过70页，需要检查图表尺寸')
    writer.close()
    document = pymupdf.open(pdf_temp)
    for i, sheet in enumerate(document):
        sheet.insert_text((34, 20), 'STOCK SELECTION | HISTORICAL AS-OF STUDY | 2026-10-13', fontsize=6.5, color=(.3, .4, .48))
        sheet.insert_text((page.width - 76, page.height - 17), f'{i+1} / {len(document)}', fontsize=8, color=(.3, .4, .48))
    document.set_metadata({'title': '板块通过后的个股筛选：详细数据与图表', 'author': 'alpha_lab', 'subject': '历史时点回放；非账户回测'})
    document.save(OUT / 'REPORT.pdf', garbage=4, deflate=True)
    extracted = [p.get_text() for p in document]
    image_pages = [i + 1 for i, p in enumerate(document) if p.get_images()]
    image_count = len({im[0] for p in document for im in p.get_images()})
    assert image_count >= 12, (image_pages, image_count)
    assert all(len(t.strip()) > 40 for t in extracted)
    assert '105' in ''.join(extracted) and '193' in ''.join(extracted)
    (OUT / 'pdf_text.txt').write_text('\n\n'.join(extracted))
    thumbw, thumbh, cols = 250, 355, 4
    canvas = Image.new('RGB', (cols * thumbw, ((len(document) + cols - 1) // cols) * thumbh), '#dce3e9')
    for i, sheet in enumerate(document):
        pix = sheet.get_pixmap(matrix=pymupdf.Matrix(.42, .42), alpha=False)
        image = Image.frombytes('RGB', [pix.width, pix.height], pix.samples); image.thumbnail((thumbw - 8, thumbh - 8))
        canvas.paste(image, ((i % cols) * thumbw + 4, (i // cols) * thumbh + 4))
    canvas.save(OUT / 'pdf_preview.png')
    document.close()
    # 工作PDF保留为排版来源，不放入交付包。
    audit = dict(status='complete', report_pages=pages, figure_groups=12, pdf_image_pages=image_pages,
                 excel_sheets=d['sheets'], source_artifacts_verified=d['verified'], path_endpoints_verified=d['path_endpoint_checks'],
                 source_run=str(RUN.relative_to(ROOT)), script_sha256=digest(__file__), account_backtest=False,
                 artifacts={str(p.relative_to(OUT)): digest(p) for p in OUT.rglob('*') if p.is_file() and p.name not in ['detail_manifest.json','stock_selection_detail.zip','layout_work.pdf']})
    (OUT / 'detail_manifest.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2))
    with zipfile.ZipFile(OUT / 'stock_selection_detail.zip', 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for path in OUT.rglob('*'):
            if path.is_file() and path.name not in ['stock_selection_detail.zip', 'layout_work.pdf']:
                archive.write(path, path.relative_to(OUT))
    print(json.dumps({k: audit[k] for k in ['status', 'report_pages', 'figure_groups', 'source_artifacts_verified', 'path_endpoints_verified']}, ensure_ascii=False))


def main():
    for path in [OUT, FIG, TABLES]:
        path.mkdir(parents=True, exist_ok=True)
    font = ROOT / 'fonts/SimHei.ttf'
    font_manager.fontManager.addfont(str(font))
    plt.rcParams.update({'font.family': font_manager.FontProperties(fname=font).get_name(), 'axes.unicode_minus':False, 'font.size':10})
    d = load_data(); print('LOADED', d['verified'], flush=True)
    compute_paths(d); print('PATHS', d['path_endpoint_checks'], flush=True)
    figures(d); print('FIGURES', 12, flush=True)
    export_tables(d); print('TABLES', len(d['sheets']), flush=True)
    text = write_report(d)
    render_and_package(d, text)
    print('REPORT', OUT / 'REPORT.pdf', flush=True)


if __name__ == '__main__':
    main()
